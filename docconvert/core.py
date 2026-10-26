"""Conversion logic.

Every reader turns a file into a list of Block. Most blocks are text
("p" or "h1".."h6") and may carry inline runs (bold/italic) and an alignment.
Readers that understand richer content also emit "img" blocks (embedded
pictures), "table" blocks (rows of cell text) and a leading "page" block
(page size and margins). Every writer turns that list back into a file and
simply skips what it cannot represent. That keeps the matrix small: adding a
format means one reader and one writer.

When LibreOffice (``soffice``) is on PATH, docx -> pdf goes through it for
exact layout fidelity. Otherwise the pure-Python path is used, which keeps the
text, pictures, tables and basic formatting but not every typographic detail.
"""

from __future__ import annotations

import base64
import html
import re
import shutil
import subprocess
import tempfile
import warnings
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = ("docx", "pdf", "md", "txt", "html")

EMU_PER_INCH = 914400


class FidelityWarning(UserWarning):
    """Raised when a conversion falls back to a path that loses layout detail."""


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class Block:
    kind: str  # "p", "h1".."h6", "img", "table" or "page"
    text: str = ""
    runs: list[Run] | None = None  # inline formatting for p/h blocks
    align: str | None = None  # "center" / "right" / "justify" (None = left)
    data: bytes | None = None  # img: raw image bytes
    mime: str | None = None  # img: e.g. "image/jpeg"
    width: float | None = None  # img: display width in inches
    height: float | None = None  # img: display height in inches
    rows: list[list[str]] | None = None  # table: cell text
    cell_blocks: list[list[list["Block"]]] | None = None  # table: rich cell content
    page: tuple[float, float] | None = None  # page: (width, height) in inches
    margins: tuple[float, float, float, float] | None = None  # page: left, right, top, bottom

    @property
    def is_text(self) -> bool:
        return self.kind == "p" or self.kind.startswith("h")


def _ext(path: Path) -> str:
    return path.suffix.lower().lstrip(".")


def _runs(block: Block) -> list[Run]:
    return block.runs if block.runs else [Run(block.text)]


# ---------------------------------------------------------------- readers


def read_txt(path: Path) -> list[Block]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    return [Block("p", chunk.strip()) for chunk in re.split(r"\n\s*\n", raw) if chunk.strip()]


def read_md(path: Path) -> list[Block]:
    blocks: list[Block] = []
    raw = path.read_text(encoding="utf-8", errors="replace")
    for chunk in re.split(r"\n\s*\n", raw):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", chunk, re.S)
        if m:
            blocks.append(Block(f"h{len(m.group(1))}", m.group(2).strip()))
        else:
            blocks.append(Block("p", chunk))
    return blocks


def read_html(path: Path) -> list[Block]:
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.blocks: list[Block] = []
            self.kind: str | None = None
            self.buf: list[str] = []

        def handle_starttag(self, tag, attrs):
            if tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "li"):
                self.kind = "p" if tag == "li" else tag
                self.buf = []
            elif tag == "br" and self.kind:
                self.buf.append("\n")

        def handle_endtag(self, tag):
            if self.kind and tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "li"):
                text = "".join(self.buf).strip()
                if text:
                    self.blocks.append(Block(self.kind, text))
                self.kind = None

        def handle_data(self, data):
            if self.kind:
                self.buf.append(data)

    parser = P()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return parser.blocks


def _merge_runs(runs: list[Run]) -> list[Run]:
    """Join neighbouring runs with identical formatting (Word splits phrases freely)."""
    merged: list[Run] = []
    for r in runs:
        if merged and (merged[-1].bold, merged[-1].italic) == (r.bold, r.italic):
            merged[-1].text += r.text
        else:
            merged.append(Run(r.text, r.bold, r.italic))
    return merged


def _docx_images(para_el, part) -> list[Block]:
    """Pictures embedded in one paragraph element, in document order."""
    from docx.oxml.ns import qn

    out: list[Block] = []
    seen: set[int] = set()
    for drawing in para_el.iter(qn("w:drawing")):
        if id(drawing) in seen:
            continue
        seen.add(id(drawing))
        extent = next(drawing.iter(qn("wp:extent")), None)
        width = height = None
        if extent is not None:
            width = int(extent.get("cx", 0)) / EMU_PER_INCH or None
            height = int(extent.get("cy", 0)) / EMU_PER_INCH or None
        blips = list(drawing.iter(qn("a:blip")))
        for blip in blips:
            rid = blip.get(qn("r:embed"))
            rel = part.rels.get(rid) if rid else None
            if rel is None or "image" not in rel.reltype:
                continue
            img = rel.target_part
            # the drawing's extent describes the whole drawing; only trust it for a single picture
            w, h = (width, height) if len(blips) == 1 else (None, None)
            out.append(Block("img", data=img.blob, mime=getattr(img, "content_type", None), width=w, height=h))
    return out


def _para_alignment(para) -> str | None:
    from docx.oxml.ns import qn

    jc = para._p.find(qn("w:pPr") + "/" + qn("w:jc"))
    if jc is None:
        return None
    val = jc.get(qn("w:val"))
    return {"center": "center", "right": "right", "end": "right", "both": "justify",
            "distribute": "justify"}.get(val)


def _para_runs(para) -> list[Run]:
    """All runs in document order, including those inside hyperlinks."""
    runs: list[Run] = []
    for item in para.iter_inner_content():
        for r in getattr(item, "runs", None) or [item]:
            if r.text:
                runs.append(Run(r.text, bool(r.bold), bool(r.italic)))
    return runs


def _docx_paragraph(para, part) -> list[Block]:
    """Text block (with runs) and/or image blocks for one python-docx Paragraph."""
    style = (para.style.name or "").lower()
    m = re.match(r"heading (\d)", style)
    kind = f"h{m.group(1)}" if m else ("h1" if style == "title" else "p")
    align = _para_alignment(para)

    blocks: list[Block] = []
    text_runs: list[Run] = []

    def flush_text() -> None:
        nonlocal text_runs
        text_runs = _merge_runs(text_runs)
        if text_runs and "".join(r.text for r in text_runs).strip():
            text_runs[0].text = text_runs[0].text.lstrip()
            text_runs[-1].text = text_runs[-1].text.rstrip()
            blocks.append(Block(kind, "".join(r.text for r in text_runs), runs=text_runs, align=align))
        text_runs = []

    # Walk runs and drawings together. This keeps an inline picture between
    # surrounding text instead of appending every picture after the paragraph.
    from docx.oxml.ns import qn
    from docx.text.run import Run as DocxRun

    for element in para._p.iter():
        if element.tag != qn("w:r"):
            continue
        run = DocxRun(element, para)
        for child in element:
            if child.tag == qn("w:t"):
                text_runs.append(Run(child.text or "", bool(run.bold), bool(run.italic)))
            elif child.tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
                text_runs.append(Run("\t" if child.tag == qn("w:tab") else "\n", bool(run.bold), bool(run.italic)))
            elif child.tag == qn("w:drawing"):
                flush_text()
                for img in _docx_images(child, part):
                    img.align = align
                    blocks.append(img)
    flush_text()
    return blocks


def read_docx(path: Path) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    blocks: list[Block] = []

    sec = doc.sections[0] if doc.sections else None
    if sec is not None and sec.page_width and sec.page_height:
        blocks.append(
            Block(
                "page",
                page=(sec.page_width.inches, sec.page_height.inches),
                margins=(
                    sec.left_margin.inches if sec.left_margin is not None else 1.0,
                    sec.right_margin.inches if sec.right_margin is not None else 1.0,
                    sec.top_margin.inches if sec.top_margin is not None else 1.0,
                    sec.bottom_margin.inches if sec.bottom_margin is not None else 1.0,
                ),
            )
        )

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            blocks.extend(_docx_paragraph(Paragraph(child, doc), doc.part))
        elif tag == "tbl":
            rows = []
            cell_blocks = []
            for row in Table(child, doc).rows:
                cells = []
                row_blocks = []
                prev_tc = None
                for cell in row.cells:
                    if cell._tc is prev_tc:  # horizontally merged cell, already recorded
                        continue
                    prev_tc = cell._tc
                    content = []
                    for p in cell.paragraphs:
                        content.extend(_docx_paragraph(p, doc.part))
                    cells.append("\n".join(b.text.strip() for b in content if b.is_text and b.text.strip()))
                    row_blocks.append(content)
                rows.append(cells)
                cell_blocks.append(row_blocks)
            if rows and any(rows):
                blocks.append(Block("table", rows=rows, cell_blocks=cell_blocks))
    return blocks


def read_pdf(path: Path) -> list[Block]:
    from pypdf import PdfReader

    text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return [Block("p", chunk.strip()) for chunk in re.split(r"\n\s*\n", text) if chunk.strip()]


# ---------------------------------------------------------------- writers


def _skip_image(b: Block) -> None:
    warnings.warn(
        f"skipped a picture of type {b.mime or 'unknown'} that the built-in renderer cannot embed",
        FidelityWarning,
        stacklevel=3,
    )


def write_txt(blocks: list[Block], path: Path) -> None:
    out = []
    for b in blocks:
        if b.is_text:
            out.append(b.text)
        elif b.kind == "table" and b.rows:
            out.append("\n".join("\t".join(c.replace("\n", " ") for c in row) for row in b.rows))
    path.write_text("\n\n".join(out) + "\n", encoding="utf-8")


_IMG_EXT = {"image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif",
            "image/bmp": "bmp", "image/tiff": "tif", "image/webp": "webp", "image/svg+xml": "svg",
            "image/x-emf": "emf", "image/x-wmf": "wmf"}


def _md_inline(block: Block) -> str:
    parts = []
    for r in _runs(block):
        t = r.text
        if not t.strip():
            parts.append(t)
            continue
        lead, body, trail = re.match(r"^(\s*)(.*?)(\s*)$", t, re.S).groups()
        if r.bold and r.italic:
            body = f"***{body}***"
        elif r.bold:
            body = f"**{body}**"
        elif r.italic:
            body = f"*{body}*"
        parts.append(lead + body + trail)
    return "".join(parts)


def write_md(blocks: list[Block], path: Path) -> None:
    out = []
    n_img = 0

    def cell_text(cell: list[Block]) -> str:
        nonlocal n_img
        parts = []
        for b in cell:
            if b.is_text:
                parts.append(_md_inline(b))
            elif b.kind == "img" and b.data:
                n_img += 1
                ext = _IMG_EXT.get(b.mime or "", "bin")
                img_path = path.with_name(f"{path.stem}_img{n_img}.{ext}")
                img_path.write_bytes(b.data)
                parts.append(f"![image {n_img}]({img_path.name})")
        return " ".join(parts).replace("\n", " ")

    for b in blocks:
        if b.kind.startswith("h"):
            out.append("#" * int(b.kind[1]) + " " + b.text)
        elif b.kind == "p":
            out.append(_md_inline(b))
        elif b.kind == "table" and b.rows and max(len(r) for r in b.rows):
            width = max(len(r) for r in b.rows)
            rows = [r + [""] * (width - len(r)) for r in b.rows]
            if b.cell_blocks:
                rows = [[cell_text(c) for c in row] for row in b.cell_blocks]
                rows = [r + [""] * (width - len(r)) for r in rows]
            esc = lambda c: c.replace("\n", " ").replace("|", "\\|")  # noqa: E731
            lines = ["| " + " | ".join(esc(c) for c in rows[0]) + " |", "|" + " --- |" * width]
            lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]
            out.append("\n".join(lines))
        elif b.kind == "img" and b.data:
            n_img += 1
            ext = _IMG_EXT.get(b.mime or "", "bin")
            img_path = path.with_name(f"{path.stem}_img{n_img}.{ext}")
            img_path.write_bytes(b.data)
            out.append(f"![image {n_img}]({img_path.name})")
    path.write_text("\n\n".join(out) + "\n", encoding="utf-8")


def _html_inline(block: Block) -> str:
    parts = []
    for r in _runs(block):
        t = html.escape(r.text).replace("\n", "<br>")
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        parts.append(t)
    return "".join(parts)


def write_html(blocks: list[Block], path: Path) -> None:
    body = []

    def cell_html(cell: list[Block]) -> str:
        parts = []
        for b in cell:
            if b.is_text:
                parts.append(_html_inline(b))
            elif b.kind == "img" and b.data:
                uri = f"data:{b.mime or 'image/png'};base64,{base64.b64encode(b.data).decode('ascii')}"
                w = f' width="{int(b.width * 96)}"' if b.width else ""
                parts.append(f'<img src="{uri}"{w} alt="">')
        return "<br>".join(parts)

    for b in blocks:
        style = f' style="text-align:{b.align}"' if b.align else ""
        if b.is_text:
            body.append(f"<{b.kind}{style}>{_html_inline(b)}</{b.kind}>")
        elif b.kind == "table" and b.rows:
            cell_rows = b.cell_blocks or [[[Block("p", c)] for c in row] for row in b.rows]
            trs = "".join(
                "<tr>" + "".join(f"<td>{cell_html(c)}</td>" for c in row) + "</tr>"
                for row in cell_rows
            )
            body.append(f'<table border="1" cellpadding="4" cellspacing="0">{trs}</table>')
        elif b.kind == "img" and b.data:
            uri = f"data:{b.mime or 'image/png'};base64,{base64.b64encode(b.data).decode('ascii')}"
            w = f' width="{int(b.width * 96)}"' if b.width else ""
            body.append(f'<p{style}><img src="{uri}"{w} alt=""></p>')
    path.write_text(
        "<!doctype html>\n<html><head><meta charset=\"utf-8\">"
        f"<title>{html.escape(path.stem)}</title></head>\n<body>\n" + "\n".join(body) + "\n</body></html>\n",
        encoding="utf-8",
    )


def write_docx(blocks: list[Block], path: Path) -> None:
    import io

    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    align_map = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                 "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    doc = docx.Document()
    for b in blocks:
        if b.kind == "page" and b.page:
            sec = doc.sections[0]
            sec.page_width, sec.page_height = Inches(b.page[0]), Inches(b.page[1])
            if b.margins:
                sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = map(Inches, b.margins)
        elif b.is_text:
            para = doc.add_heading("", level=int(b.kind[1])) if b.kind.startswith("h") else doc.add_paragraph()
            for r in _runs(b):
                run = para.add_run(r.text)
                run.bold, run.italic = r.bold or None, r.italic or None
            if b.align in align_map:
                para.alignment = align_map[b.align]
        elif b.kind == "table" and b.rows and max(len(r) for r in b.rows):
            width = max(len(r) for r in b.rows)
            table = doc.add_table(rows=len(b.rows), cols=width)
            table.style = "Table Grid"
            for i, row in enumerate(b.rows):
                for j, cell in enumerate(row):
                    target = table.cell(i, j)
                    content = b.cell_blocks[i][j] if b.cell_blocks and i < len(b.cell_blocks) and j < len(b.cell_blocks[i]) else [Block("p", cell)]
                    target.text = ""
                    para = target.paragraphs[0]
                    for item in content:
                        if item.is_text:
                            for r in _runs(item):
                                run = para.add_run(r.text)
                                run.bold, run.italic = r.bold or None, r.italic or None
                        elif item.kind == "img" and item.data:
                            try:
                                para.add_run().add_picture(io.BytesIO(item.data), width=Inches(item.width) if item.width else None)
                            except Exception:
                                _skip_image(item)
        elif b.kind == "img" and b.data:
            try:
                doc.add_picture(io.BytesIO(b.data), width=Inches(b.width) if b.width else None)
            except Exception:  # noqa: BLE001 - e.g. EMF/WMF/SVG that python-docx cannot embed
                _skip_image(b)
                continue
            if b.align in align_map:
                doc.paragraphs[-1].alignment = align_map[b.align]
    doc.save(str(path))


def write_pdf(blocks: list[Block], path: Path) -> None:
    import io

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    page = blocks[0] if blocks and blocks[0].kind == "page" else None
    pagesize = tuple(v * inch for v in page.page) if page and page.page else A4
    margins = page.margins if page and page.margins else (1.0, 1.0, 1.0, 1.0)
    left, right, top, bottom = (m * inch for m in margins)
    frame_w = pagesize[0] - left - right
    frame_h = pagesize[1] - top - bottom
    if frame_w < inch or frame_h < inch:  # nonsensical geometry in the source: fall back to A4
        pagesize, (left, right, top, bottom) = A4, (inch,) * 4
        frame_w, frame_h = pagesize[0] - 2 * inch, pagesize[1] - 2 * inch

    styles = getSampleStyleSheet()
    align_map = {"center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}

    def style_for(b: Block) -> ParagraphStyle:
        base = styles["Heading" + b.kind[1]] if b.kind.startswith("h") else styles["BodyText"]
        if b.align in align_map:
            return ParagraphStyle(f"{base.name}-{b.align}", parent=base, alignment=align_map[b.align])
        return base

    def markup(b: Block) -> str:
        parts = []
        for r in _runs(b):
            t = html.escape(r.text).replace("\n", "<br/>")
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            parts.append(t)
        return "".join(parts)

    def cell_flowables(cell: list[Block], cell_style: ParagraphStyle) -> list:
        flowables = []
        for item in cell:
            if item.is_text:
                flowables.append(Paragraph(markup(item), cell_style))
            elif item.kind == "img" and item.data:
                try:
                    px_w, px_h = ImageReader(io.BytesIO(item.data)).getSize()
                    if not px_w or not px_h:
                        raise ValueError("empty image")
                    w = item.width * inch if item.width else min(frame_w / 2, px_w * 72 / 96)
                    h = item.height * inch if (item.height and item.width) else w * px_h / px_w
                    scale = min(1.0, (frame_w / 2) / w, (frame_h - 12) / h)
                    flowables.append(Image(io.BytesIO(item.data), width=w * scale, height=h * scale))
                except Exception:
                    _skip_image(item)
        return flowables or [Paragraph("", cell_style)]

    story = []
    for b in blocks:
        if b.is_text:
            story.append(Paragraph(markup(b), style_for(b)))
            story.append(Spacer(1, 6))
        elif b.kind == "table" and b.rows and max(len(r) for r in b.rows):
            width = max(len(r) for r in b.rows)
            cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=9, leading=11)
            cell_rows = b.cell_blocks or [[[Block("p", c)] for c in row] for row in b.rows]
            data = [
                [cell_flowables(c, cell_style) for c in r + [[]] * (width - len(r))]
                for r in cell_rows
            ]
            t = Table(data, colWidths=[frame_w / width] * width, repeatRows=0)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
        elif b.kind == "img" and b.data:
            try:
                px_w, px_h = ImageReader(io.BytesIO(b.data)).getSize()
            except Exception:  # noqa: BLE001 - e.g. EMF/WMF/SVG that Pillow cannot decode
                _skip_image(b)
                continue
            if not px_w or not px_h:
                _skip_image(b)
                continue
            w = b.width * inch if b.width else min(frame_w, px_w * 72 / 96)
            h = b.height * inch if (b.height and b.width) else w * px_h / px_w
            scale = min(1.0, frame_w / w, (frame_h - 12) / h)  # never overflow the page frame
            img = Image(io.BytesIO(b.data), width=w * scale, height=h * scale)
            img.hAlign = {"center": "CENTER", "right": "RIGHT"}.get(b.align or "", "LEFT")
            story.append(img)
            story.append(Spacer(1, 8))

    SimpleDocTemplate(
        str(path), pagesize=pagesize, title=path.stem,
        leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
    ).build(story)


READERS = {"txt": read_txt, "md": read_md, "html": read_html, "docx": read_docx, "pdf": read_pdf}
WRITERS = {"txt": write_txt, "md": write_md, "html": write_html, "docx": write_docx, "pdf": write_pdf}

# ---------------------------------------------------------------- libreoffice


def _soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def _convert_with_soffice(src: Path, dst: Path, exe: str) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [exe, "--headless", "--convert-to", _ext(dst), "--outdir", tmp, str(src)],
            check=True,
            capture_output=True,
            timeout=120,
        )
        produced = Path(tmp) / (src.stem + "." + _ext(dst))
        if not produced.exists():
            raise RuntimeError("LibreOffice did not produce an output file")
        shutil.move(str(produced), str(dst))


# ---------------------------------------------------------------- public API


def convert(src: str | Path, dst: str | Path, *, use_soffice: bool = True) -> Path:
    """Convert ``src`` to ``dst``; the formats are taken from the extensions."""
    src, dst = Path(src), Path(dst)
    sfmt, dfmt = _ext(src), _ext(dst)
    if sfmt not in READERS:
        raise ValueError(f"unsupported input format: .{sfmt} (supported: {', '.join(SUPPORTED)})")
    if dfmt not in WRITERS:
        raise ValueError(f"unsupported output format: .{dfmt} (supported: {', '.join(SUPPORTED)})")
    if not src.is_file():
        raise FileNotFoundError(f"file not found: {src}")
    if sfmt == dfmt:
        raise ValueError("input and output formats are the same")

    dst.parent.mkdir(parents=True, exist_ok=True)

    if sfmt == "docx" and dfmt == "pdf":
        exe = _soffice() if use_soffice else None
        if exe:
            _convert_with_soffice(src, dst, exe)
            return dst
        if use_soffice:
            warnings.warn(
                "LibreOffice not found: using the built-in docx -> pdf renderer. Text, pictures, "
                "tables and basic formatting are kept, but the page layout may differ from Word. "
                "Install LibreOffice for an exact rendering.",
                FidelityWarning,
                stacklevel=2,
            )

    WRITERS[dfmt](READERS[sfmt](src), dst)
    return dst
