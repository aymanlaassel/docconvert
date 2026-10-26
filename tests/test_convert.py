import itertools
from pathlib import Path

import pytest

from docconvert import SUPPORTED, convert
from docconvert.core import READERS

SAMPLE_MD = "# Title\n\nFirst paragraph with some text.\n\n## Section\n\nSecond paragraph.\n"


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    p = tmp_path / "sample.md"
    p.write_text(SAMPLE_MD, encoding="utf-8")
    return p


@pytest.mark.parametrize("src_fmt,dst_fmt", [p for p in itertools.permutations(SUPPORTED, 2)])
def test_every_pair_produces_text(tmp_path: Path, sample_md: Path, src_fmt: str, dst_fmt: str):
    src = sample_md if src_fmt == "md" else convert(sample_md, tmp_path / f"src.{src_fmt}", use_soffice=False)
    dst = convert(src, tmp_path / f"out.{dst_fmt}", use_soffice=False)
    assert dst.is_file() and dst.stat().st_size > 0
    text = " ".join(b.text for b in READERS[dst_fmt](dst))
    assert "First paragraph" in text
    assert "Second paragraph" in text


def test_headings_survive_md_docx_md(tmp_path: Path, sample_md: Path):
    docx = convert(sample_md, tmp_path / "a.docx")
    back = convert(docx, tmp_path / "b.md")
    assert back.read_text().startswith("# Title")
    assert "## Section" in back.read_text()


def test_rejects_unknown_format(tmp_path: Path, sample_md: Path):
    with pytest.raises(ValueError):
        convert(sample_md, tmp_path / "x.xyz")


def test_missing_input(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        convert(tmp_path / "nope.md", tmp_path / "x.pdf")


def test_cli_batch(tmp_path: Path, sample_md: Path):
    from docconvert.cli import main

    assert main([str(sample_md), "--to", "txt", "--no-soffice"]) == 0
    assert (tmp_path / "sample.txt").is_file()


# ---------------------------------------------------------------- rich docx content


@pytest.fixture
def rich_docx(tmp_path: Path) -> Path:
    """A docx with a picture, a table, alignment and bold/italic runs."""
    import io

    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 100), "red").save(buf, format="PNG")

    d = docx.Document()
    d.add_heading("Rich Title", level=1)
    p = d.add_paragraph()
    p.add_run("Bold part").bold = True
    p.add_run(" and ")
    p.add_run("italic part").italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Key", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "Course", "COT 3541"
    d.add_paragraph("Picture below:")
    d.add_picture(io.BytesIO(buf.getvalue()), width=Inches(2))
    out = tmp_path / "rich.docx"
    d.save(str(out))
    return out


def test_docx_reader_keeps_images_tables_and_runs(rich_docx: Path):
    from docconvert.core import read_docx

    blocks = read_docx(rich_docx)
    kinds = [b.kind for b in blocks]
    assert kinds[0] == "page" and blocks[0].page is not None
    assert "img" in kinds and "table" in kinds
    img = next(b for b in blocks if b.kind == "img")
    assert img.data and img.data[:4] == b"\x89PNG" and abs(img.width - 2.0) < 0.01
    table = next(b for b in blocks if b.kind == "table")
    assert table.rows == [["Key", "Value"], ["Course", "COT 3541"]]
    para = next(b for b in blocks if b.kind == "p" and b.runs)
    assert para.align == "center"
    assert [(r.text, r.bold, r.italic) for r in para.runs] == [
        ("Bold part", True, False), (" and ", False, False), ("italic part", False, True)]


def test_inline_picture_stays_between_surrounding_text(tmp_path: Path):
    import io

    import docx
    from PIL import Image
    from docconvert.core import read_docx

    buf = io.BytesIO()
    Image.new("RGB", (20, 20), "red").save(buf, format="PNG")
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("before")
    p.add_run().add_picture(io.BytesIO(buf.getvalue()))
    p.add_run("after")
    src = tmp_path / "inline.docx"
    d.save(str(src))

    blocks = read_docx(src)
    assert [b.kind for b in blocks[1:]] == ["p", "img", "p"]
    assert [b.text for b in blocks[1:]] == ["before", "", "after"]


def test_docx_to_pdf_fallback_keeps_images_and_tables(tmp_path: Path, rich_docx: Path):
    from pypdf import PdfReader

    pdf = convert(rich_docx, tmp_path / "rich.pdf", use_soffice=False)
    reader = PdfReader(str(pdf))
    assert sum(len(p.images) for p in reader.pages) == 1
    text = " ".join(p.extract_text() for p in reader.pages)
    assert "COT 3541" in text and "Bold part" in text


def test_warns_when_soffice_missing(tmp_path: Path, rich_docx: Path, monkeypatch):
    from docconvert import core

    monkeypatch.setattr(core, "_soffice", lambda: None)
    with pytest.warns(core.FidelityWarning):
        convert(rich_docx, tmp_path / "rich.pdf", use_soffice=True)


def test_no_warning_when_soffice_explicitly_disabled(tmp_path: Path, rich_docx: Path):
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        convert(rich_docx, tmp_path / "rich.pdf", use_soffice=False)


@pytest.mark.parametrize("dst_fmt", ["html", "md"])
def test_rich_docx_to_other_formats_keep_image(tmp_path: Path, rich_docx: Path, dst_fmt: str):
    out = convert(rich_docx, tmp_path / f"rich.{dst_fmt}")
    if dst_fmt == "html":
        assert "data:image/png;base64," in out.read_text() and "<table" in out.read_text()
    else:
        assert "![image 1](rich_img1.png)" in out.read_text() and (tmp_path / "rich_img1.png").is_file()
        assert "| Course | COT 3541 |" in out.read_text()


def test_docx_writer_keeps_image_and_table(tmp_path: Path, rich_docx: Path):
    import docx

    from docconvert.core import read_docx, write_docx

    out = tmp_path / "copy.docx"
    write_docx(read_docx(rich_docx), out)
    d = docx.Document(str(out))
    assert any("image" in r.reltype for r in d.part.rels.values())
    assert d.tables and d.tables[0].cell(1, 1).text == "COT 3541"


def test_picture_in_table_cell_stays_in_cell(tmp_path: Path):
    import io

    import docx
    from PIL import Image
    from docconvert.core import convert, read_docx

    buf = io.BytesIO()
    Image.new("RGB", (20, 20), "blue").save(buf, format="PNG")
    d = docx.Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("caption")
    cell.add_paragraph().add_run().add_picture(io.BytesIO(buf.getvalue()))
    src = tmp_path / "cell-image.docx"
    d.save(str(src))

    blocks = read_docx(src)
    table = next(b for b in blocks if b.kind == "table")
    assert table.cell_blocks and any(b.kind == "img" for b in table.cell_blocks[0][0])
    out = tmp_path / "cell-image.html"
    convert(src, out)
    markup = out.read_text()
    assert "<table" in markup and markup.index("<table") < markup.index("data:image/png;base64,")


# ---------------------------------------------------------------- awkward docx input


def _docx_with_body_xml(tmp_path: Path, name: str, body_xml: str) -> Path:
    """Build a docx from python-docx and splice raw body XML in (for things the API can't write)."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    d.element.body.insert(0, parse_xml(f"<w:p {nsdecls('w', 'r')}>{body_xml}</w:p>"))
    out = tmp_path / name
    d.save(str(out))
    return out


def test_hyperlink_text_is_kept(tmp_path: Path):
    from docconvert.core import read_docx

    src = _docx_with_body_xml(
        tmp_path, "link.docx",
        '<w:r><w:t xml:space="preserve">See </w:t></w:r>'
        '<w:hyperlink r:id="rIdX"><w:r><w:t>the link</w:t></w:r></w:hyperlink>'
        '<w:r><w:t xml:space="preserve"> now</w:t></w:r>',
    )
    assert read_docx(src)[1].text == "See the link now"


def test_nonstandard_alignment_does_not_crash(tmp_path: Path):
    from docconvert.core import read_docx

    src = _docx_with_body_xml(
        tmp_path, "start.docx", '<w:pPr><w:jc w:val="start"/></w:pPr><w:r><w:t>hello</w:t></w:r>'
    )
    b = read_docx(src)[1]
    assert b.text == "hello" and b.align is None


def test_split_bold_runs_are_merged(tmp_path: Path):
    from docconvert.core import read_docx, write_md

    src = _docx_with_body_xml(
        tmp_path, "split.docx",
        '<w:r><w:rPr><w:b/></w:rPr><w:t>Bold</w:t></w:r>'
        '<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve"> Part</w:t></w:r>',
    )
    blocks = read_docx(src)
    assert len(blocks[1].runs) == 1
    out = tmp_path / "split.md"
    write_md(blocks, out)
    assert out.read_text().strip() == "**Bold Part**"


def test_merged_table_cells_not_repeated(tmp_path: Path):
    import docx

    from docconvert.core import read_docx

    d = docx.Document()
    t = d.add_table(rows=1, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged"
    t.cell(0, 2).text = "C"
    src = tmp_path / "merged.docx"
    d.save(str(src))
    table = next(b for b in read_docx(src) if b.kind == "table")
    assert table.rows == [["Merged", "C"]]


def test_unembeddable_picture_is_skipped_with_warning(tmp_path: Path):
    from pypdf import PdfReader

    from docconvert.core import Block, FidelityWarning, write_docx, write_pdf

    blocks = [Block("p", "before"), Block("img", data=b"\x01\x00\x00\x00not-an-image", mime="image/x-emf"),
              Block("p", "after")]
    with pytest.warns(FidelityWarning):
        write_pdf(blocks, tmp_path / "emf.pdf")
    assert "after" in PdfReader(str(tmp_path / "emf.pdf")).pages[0].extract_text()
    with pytest.warns(FidelityWarning):
        write_docx(blocks, tmp_path / "emf.docx")


def test_empty_table_rows_do_not_crash(tmp_path: Path):
    from docconvert.core import Block, write_html, write_md, write_pdf

    blocks = [Block("table", rows=[[]]), Block("p", "ok")]
    write_pdf(blocks, tmp_path / "t.pdf")
    write_md(blocks, tmp_path / "t.md")
    write_html(blocks, tmp_path / "t.html")


def test_absurd_margins_fall_back_to_a4(tmp_path: Path):
    from pypdf import PdfReader

    from docconvert.core import Block, write_pdf

    blocks = [Block("page", page=(3.0, 3.0), margins=(2.0, 2.0, 2.0, 2.0)), Block("p", "tiny page")]
    write_pdf(blocks, tmp_path / "m.pdf")
    assert "tiny page" in PdfReader(str(tmp_path / "m.pdf")).pages[0].extract_text()
